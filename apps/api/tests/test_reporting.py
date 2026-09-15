from io import BytesIO
from decimal import Decimal

from openpyxl import load_workbook

from conftest import login


def relogin(client, email):
    client.cookies.clear()
    return login(client, email)


def test_complete_reporting_drilldowns_exports_and_role_scope(client):
    login(client, "alex@atplcrm.local")
    response = client.get("/api/v1/reports/analytics/?period=all&group_by=country")
    assert response.status_code == 200, response.text
    report = response.json()
    assert {
        "pipeline", "forecast_month", "forecast_quarter", "lead_funnel",
        "lead_by_source", "lead_by_user", "blockers_by_type",
        "blockers_by_owner", "outcomes", "loss_reasons", "value_erosion",
        "movement", "performance", "milestones", "records",
    } <= report.keys()
    assert len(report["milestones"]) == 7
    assert [row["stage"] for row in report["lead_funnel"]] == [
        "Created", "Worked", "Engaged", "Validated", "Disqualified", "Converted"
    ]
    assert {row["key"] for row in report["movement"]} == {
        "entered", "advanced", "regressed", "closed", "slipped"
    }
    assert Decimal(report["pipeline_total_usd"]) == sum(
        (Decimal(row["net_value_usd"]) for row in report["pipeline"]), Decimal("0")
    )
    known_ids = {row["id"] for row in report["records"]}
    for collection in (
        report["pipeline"], report["forecast_month"], report["lead_funnel"],
        report["blockers_by_type"], report["outcomes"], report["movement"],
        report["milestones"],
    ):
        assert all(set(row["record_ids"]) <= known_ids for row in collection)
    stages = {row["id"]: row["stage"] for row in report["records"]}
    assert all(stages[identifier] != "hold" for row in report["forecast_month"] for identifier in row["record_ids"])

    home = client.get("/api/v1/reports/home/").json()
    assert home["role"] == "Head of Sales"
    assert len(home["cards"]) == 4
    assert all(card["route"] and isinstance(card["record_ids"], list) for card in home["cards"])

    csv_export = client.get("/api/v1/reports/export/?report=outcomes&format=csv&period=all")
    assert csv_export.status_code == 200
    assert csv_export.content.startswith(b"\xef\xbb\xbf")
    workbook = client.get("/api/v1/reports/export/?report=performance&format=xlsx&period=all")
    assert workbook.status_code == 200
    assert load_workbook(BytesIO(workbook.content)).active.max_row >= 2
    list_csv = client.get("/api/v1/productivity/lists/opportunities/?export_format=csv")
    assert list_csv.status_code == 200 and b"expected_close_date" in list_csv.content
    list_xlsx = client.get("/api/v1/productivity/lists/contacts/?export_format=xlsx")
    assert list_xlsx.status_code == 200
    assert load_workbook(BytesIO(list_xlsx.content)).active.max_row >= 2

    relogin(client, "maya@atplcrm.local")
    individual = client.get("/api/v1/reports/analytics/?period=all").json()
    assert len(individual["performance"]) == 1
    assert individual["performance"][0]["user"] == "Maya Patel"
