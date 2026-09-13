"""Tenant-scoped imports, search, duplicate review and data-quality endpoints."""
from __future__ import annotations

import csv
import base64
import binascii
import io
import re
from zipfile import BadZipFile
from collections import defaultdict
from datetime import date, datetime, timezone
from difflib import SequenceMatcher
from itertools import combinations
from uuid import UUID

from email_validator import EmailNotValidError, validate_email
from fastapi import APIRouter, Depends, Query
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, or_, select, update
from sqlalchemy.orm import Session, selectinload

from .constants import COMPANY_TYPES, CONSENT_BASES, CONTACT_SENIORITIES, ENGAGEMENT_STATUSES, MANAGEMENT
from .database import get_db
from .models import Activity, Company, Contact, DuplicateDecision, ImportJob, Lead, Opportunity, PartnerInvolvement, Pursuit, PursuitContact, User
from .references import codes as reference_codes
from .schemas import DuplicateDismissInput, ImportInput, MergeInput
from .security import current_user
from .services import audit, future_date, http_error, scoped, stamp

router = APIRouter(prefix="/api/v1/data", tags=["data tools"])

TEMPLATES = {
    "companies": ["name", "country", "owner_email", "domain", "industry", "company_type", "primary_region", "global_account_name"],
    "contacts": ["company_name", "first_name", "last_name", "country", "owner_email", "sourced_by_email", "email", "job_title", "seniority", "phone", "mobile", "linkedin_url", "city", "source_channel", "source_detail", "engagement_status", "do_not_contact", "consent_basis", "notes"],
    "leads": ["name", "company_name", "owner_email", "holder_email", "sourced_by_email", "next_action", "action_type", "action_date", "source_channel", "source_detail", "priority", "area_of_interest"],
}
REQUIRED = {
    "companies": {"name", "country", "owner_email"},
    "contacts": {"company_name", "first_name", "last_name", "country", "owner_email"},
    "leads": {"name", "company_name", "owner_email", "next_action", "action_type", "action_date", "source_channel"},
}


def normalized(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").casefold())


def import_access(user: User) -> None:
    if user.level not in MANAGEMENT:
        raise http_error(403, "Manager, Executive or Administrator access is required for imports and merges.")


def cell_text(value) -> str:
    if value is None: return ""
    if isinstance(value, datetime): return value.isoformat()
    if isinstance(value, date): return value.isoformat()
    if isinstance(value, bool): return "true" if value else "false"
    return str(value).strip()


def parse(payload: ImportInput) -> tuple[list[str], list[dict[str, str]]]:
    if payload.file_type == "xlsx":
        try:
            content = base64.b64decode(payload.file_content, validate=True)
            if len(content) > 8_000_000: raise http_error(422, {"file_content": "The Excel file must be 8 MB or smaller."})
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, None)
            headers = [cell_text(value) for value in (first or [])]
            rows = [{headers[index]: cell_text(value) for index, value in enumerate(row) if index < len(headers) and headers[index]} for row in iterator if any(value is not None and cell_text(value) for value in row)]
            workbook.close()
        except (binascii.Error, BadZipFile, InvalidFileException, OSError, ValueError, KeyError, EOFError) as exc:
            raise http_error(422, {"file_content": "Could not read this .xlsx workbook. Use a valid, unencrypted Excel file."}) from exc
    else:
        if len(payload.csv_text.encode("utf-8")) > 8_000_000: raise http_error(422, {"csv_text": "The CSV file must be 8 MB or smaller."})
        try:
            reader = csv.DictReader(io.StringIO(payload.csv_text.lstrip("\ufeff")))
            headers = [str(value).strip() if value is not None else "" for value in (reader.fieldnames or [])]
            rows = [{str(k).strip(): (v or "").strip() for k, v in row.items() if k} for row in reader if any((value or "").strip() for value in row.values())]
        except csv.Error as exc:
            raise http_error(422, {"csv_text": f"Could not read this CSV: {exc}"}) from exc
    if not headers:
        raise http_error(422, {"file": "The file needs a header row."})
    normalized_headers = [normalized(header) for header in headers]
    if "" in normalized_headers or len(set(normalized_headers)) != len(normalized_headers):
        raise http_error(422, {"file": "Every source column needs a unique, non-empty heading."})
    if len(rows) > 20_000:
        raise http_error(422, {"file": "A single import can contain at most 20,000 rows."})
    return headers, rows


def source_for(field: str, headers: list[str], mapping: dict[str, str]) -> str | None:
    requested = mapping.get(field, field)
    by_normalized = {normalized(header): header for header in headers}
    return by_normalized.get(normalized(requested))


def canonical_rows(payload: ImportInput) -> tuple[list[str], list[dict[str, str]], list[str]]:
    headers, raw = parse(payload)
    fields = TEMPLATES[payload.entity_type]
    sources = {field: source_for(field, headers, payload.mapping) for field in fields}
    missing = sorted(field for field in REQUIRED[payload.entity_type] if not sources[field])
    rows = [{field: row.get(source, "") if source else "" for field, source in sources.items()} for row in raw]
    return headers, rows, missing


def import_context(db: Session, user: User) -> dict:
    users = db.scalars(select(User).where(User.tenant_id == user.tenant_id, User.is_active.is_(True))).all()
    companies = db.scalars(scoped(db, Company, user)).all()
    contacts = db.scalars(scoped(db, Contact, user)).all()
    pursuits = db.scalars(scoped(db, Pursuit, user)).all()
    company_domains = defaultdict(list); company_blocks = defaultdict(list); company_name_matches = defaultdict(list)
    for item in companies:
        if normalized(item.domain): company_domains[normalized(item.domain)].append(item)
        company_name_matches[normalized(item.name)].append(item)
        company_blocks[(normalized(item.country), normalized(item.name)[:3])].append(item)
    contact_phones = defaultdict(list); contact_blocks = defaultdict(list)
    for item in contacts:
        phone = normalized(item.mobile or item.phone)
        if phone: contact_phones[phone].append(item)
        contact_blocks[(item.company_id, normalized(item.last_name)[:2])].append(item)
    return {
        "users": {item.email.casefold(): item for item in users},
        "companies": {item.name.casefold(): item for item in companies},
        "company_name_matches": company_name_matches,
        "contact_emails": {item.email.casefold() for item in contacts if item.email},
        "pursuits": {(item.company_id, item.name.casefold()) for item in pursuits},
        "company_domains": company_domains,
        "company_blocks": company_blocks,
        "contact_phones": contact_phones,
        "contact_blocks": contact_blocks,
        "actions": reference_codes(db, user, "actions"),
        "sources": reference_codes(db, user, "sources"),
    }


def row_errors(context: dict, entity: str, row: dict[str, str], number: int) -> list[dict]:
    errors: list[dict] = []
    for field in REQUIRED[entity]:
        if not row.get(field):
            errors.append({"row": number, "field": field, "message": "Required value is missing."})
    if errors:
        return errors
    owner = context["users"].get(row["owner_email"].casefold())
    if not owner:
        errors.append({"row": number, "field": "owner_email", "message": "No active workspace user has this email."})
    if entity == "companies":
        if row.get("company_type") and row["company_type"] not in COMPANY_TYPES:
            errors.append({"row": number, "field": "company_type", "message": "Choose a supported company type."})
    elif entity == "contacts":
        company = context["companies"].get(row["company_name"].casefold())
        if not company:
            errors.append({"row": number, "field": "company_name", "message": "Company does not exist. Import companies first."})
        if row.get("email"):
            try:
                validate_email(row["email"], check_deliverability=False)
            except EmailNotValidError:
                errors.append({"row": number, "field": "email", "message": "Enter a valid email address."})
            if row["email"].casefold() in context["contact_emails"]:
                errors.append({"row": number, "field": "email", "message": "Contact email already exists."})
        for field in ("sourced_by_email",):
            if row.get(field) and row[field].casefold() not in context["users"]:
                errors.append({"row": number, "field": field, "message": "No active workspace user has this email."})
        if row.get("source_channel"):
            if row["source_channel"] not in context["sources"]: errors.append({"row": number, "field": "source_channel", "message": "Value is not active in workspace configuration."})
        if row.get("seniority") and row["seniority"] not in CONTACT_SENIORITIES:
            errors.append({"row": number, "field": "seniority", "message": "Choose a supported seniority."})
        if row.get("engagement_status") and row["engagement_status"] not in ENGAGEMENT_STATUSES:
            errors.append({"row": number, "field": "engagement_status", "message": "Choose a supported engagement status."})
        if row.get("consent_basis") and row["consent_basis"] not in CONSENT_BASES:
            errors.append({"row": number, "field": "consent_basis", "message": "Choose a supported consent basis."})
        if row.get("do_not_contact") and row["do_not_contact"].casefold() not in {"true", "false", "yes", "no", "1", "0"}:
            errors.append({"row": number, "field": "do_not_contact", "message": "Use Yes/No or True/False."})
        if row.get("linkedin_url") and not row["linkedin_url"].casefold().startswith(("https://", "http://")):
            errors.append({"row": number, "field": "linkedin_url", "message": "Use a complete http:// or https:// URL."})
    else:
        company = context["companies"].get(row["company_name"].casefold())
        if not company:
            errors.append({"row": number, "field": "company_name", "message": "Company does not exist. Import companies first."})
        else:
            if (company.id, row["name"].casefold()) in context["pursuits"]:
                errors.append({"row": number, "field": "name", "message": "A pursuit with this name already exists at the company."})
        try:
            parsed = date.fromisoformat(row["action_date"])
            if parsed <= date.today():
                errors.append({"row": number, "field": "action_date", "message": "Use a date after today (YYYY-MM-DD)."})
        except ValueError:
            errors.append({"row": number, "field": "action_date", "message": "Use YYYY-MM-DD format."})
        for category, field in (("actions", "action_type"), ("sources", "source_channel")):
            if row[field] not in context[category]: errors.append({"row": number, "field": field, "message": "Value is not active in workspace configuration."})
        for field in ("holder_email", "sourced_by_email"):
            if row.get(field) and row[field].casefold() not in context["users"]:
                errors.append({"row": number, "field": field, "message": "No active workspace user has this email."})
    return errors


def within_file_errors(entity: str, rows: list[dict[str, str]]) -> list[dict]:
    errors: list[dict] = []
    seen: dict[tuple[str, str], int] = {}
    if entity == "companies":
        return errors
    keys = {"contacts": ("email",), "leads": ("company_name", "name")}[entity]
    for number, row in enumerate(rows, start=2):
        if entity == "leads":
            candidates = [("name", f'{normalized(row.get("company_name", ""))}:{normalized(row.get("name", ""))}')]
        else:
            candidates = [(field, normalized(row.get(field, ""))) for field in keys]
        for field, value in candidates:
            if not value: continue
            key = (field, value)
            if key in seen:
                errors.append({"row": number, "field": field, "message": f"Duplicates row {seen[key]} in this file."})
            else: seen[key] = number
    return errors


def within_file_warnings(entity: str, rows: list[dict[str, str]]) -> list[dict]:
    if entity != "companies":
        return []
    warnings: list[dict] = []
    seen: dict[tuple[str, str], int] = {}
    for number, row in enumerate(rows, start=2):
        for field in ("name", "domain"):
            value = normalized(row.get(field, ""))
            if not value:
                continue
            key = (field, value)
            if key in seen:
                warnings.append({"row": number, "field": field, "message": f"Possible duplicate of row {seen[key]}: same company {field}.", "record_id": f"file-row-{seen[key]}", "confidence": 100})
            else:
                seen[key] = number
    return warnings


def similarity(left: str, right: str) -> int:
    return round(100 * SequenceMatcher(None, normalized(left), normalized(right)).ratio())


def row_warnings(context: dict, entity: str, row: dict[str, str], number: int) -> list[dict]:
    warnings: list[dict] = []
    if entity == "companies":
        candidates = {}
        for item in context["company_name_matches"].get(normalized(row.get("name", "")), []): candidates[item.id] = item
        for item in context["company_domains"].get(normalized(row.get("domain", "")), []): candidates[item.id] = item
        block = (normalized(row.get("country", "")), normalized(row.get("name", ""))[:3])
        for item in context["company_blocks"].get(block, []): candidates[item.id] = item
        for item in candidates.values():
            reasons = []
            confidence = similarity(row.get("name", ""), item.name)
            if normalized(row.get("name", "")) == normalized(item.name): reasons.append("same company name"); confidence = 100
            if row.get("domain") and normalized(row["domain"]) == normalized(item.domain): reasons.append("same domain"); confidence = 100
            if not reasons and confidence >= 88 and row.get("country", "").casefold() == item.country.casefold(): reasons.append("similar name in the same country")
            if reasons: warnings.append({"row": number, "field": "name", "message": f'Possible duplicate of {item.name}: {", ".join(reasons)}.', "record_id": str(item.id), "confidence": confidence})
    elif entity == "contacts":
        company = context["companies"].get(row.get("company_name", "").casefold())
        if company:
            full_name = f'{row.get("first_name", "")} {row.get("last_name", "")}'
            candidates = {}
            phone = normalized(row.get("mobile") or row.get("phone", ""))
            for item in context["contact_phones"].get(phone, []): candidates[item.id] = item
            block = (company.id, normalized(row.get("last_name", ""))[:2])
            for item in context["contact_blocks"].get(block, []): candidates[item.id] = item
            for item in candidates.values():
                confidence = similarity(full_name, item.name); reasons = []
                incoming_phone = normalized(row.get("mobile") or row.get("phone", "")); existing_phone = normalized(item.mobile or item.phone)
                if incoming_phone and incoming_phone == existing_phone: reasons.append("same phone"); confidence = 100
                elif confidence >= 90: reasons.append("similar name at the same company")
                if reasons: warnings.append({"row": number, "field": "first_name", "message": f'Possible duplicate of {item.name}: {", ".join(reasons)}.', "record_id": str(item.id), "confidence": confidence})
    return warnings


@router.get("/templates/{entity_type}/")
def template(entity_type: str, user: User = Depends(current_user)):
    if entity_type not in TEMPLATES:
        raise http_error(404, "Import template not found.")
    example = {
        "companies": ["Example Industries", "United States", user.email, "example.com", "Manufacturing", "Prospect", "North America", ""],
        "contacts": ["Example Industries", "Avery", "Stone", "United States", user.email, user.email, "avery@example.com", "Operations Director", "Director", "+1 555 0100", "", "https://linkedin.com/in/avery-stone", "Austin", "LinkedIn", "Direct connection", "Not contacted", "No", "Public professional profile", "Met through professional network"],
        "leads": ["Analytics modernization", "Example Industries", user.email, user.email, user.email, "Confirm discovery call", "Call", str(date.today().replace(year=date.today().year + 1)), "LinkedIn", "Direct outreach", "Medium", "Analytics"],
    }[entity_type]
    stream = io.StringIO(); writer = csv.writer(stream); writer.writerow(TEMPLATES[entity_type]); writer.writerow(example)
    return {"entity_type": entity_type, "filename": f"atplcrm-{entity_type}-template.csv", "csv_text": stream.getvalue(), "required": sorted(REQUIRED[entity_type])}


@router.post("/imports/preview/")
def preview_import(payload: ImportInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    import_access(user)
    headers, rows, missing = canonical_rows(payload)
    context = import_context(db, user)
    errors = [{"row": 1, "field": field, "message": "Map this required field to a source column."} for field in missing]
    if not missing:
        for number, row in enumerate(rows, start=2):
            errors.extend(row_errors(context, payload.entity_type, row, number))
        errors.extend(within_file_errors(payload.entity_type, rows))
    invalid_rows = {error["row"] for error in errors if error["row"] > 1}
    warnings = [] if missing else [warning for number, row in enumerate(rows, start=2) if number not in invalid_rows for warning in row_warnings(context, payload.entity_type, row, number)]
    if not missing:
        warnings.extend(within_file_warnings(payload.entity_type, rows))
    return {"headers": headers, "suggested_mapping": {field: source_for(field, headers, payload.mapping) for field in TEMPLATES[payload.entity_type]}, "sample": rows[:5], "total_rows": len(rows), "valid_rows": max(0, len(rows) - len(invalid_rows)) if not missing else 0, "invalid_rows": len(invalid_rows), "errors": errors[:200], "truncated_errors": len(errors) > 200, "warnings": warnings[:200], "truncated_warnings": len(warnings) > 200}


@router.post("/imports/", status_code=201)
def execute_import(payload: ImportInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    import_access(user)
    _headers, rows, missing = canonical_rows(payload)
    if missing:
        raise http_error(422, {field: "Map this required field to a source column." for field in missing})
    context = import_context(db, user)
    file_errors = within_file_errors(payload.entity_type, rows)
    if file_errors: raise http_error(422, {"file": "Resolve duplicate rows in the file before importing."})
    validation = {number: row_errors(context, payload.entity_type, row, number) for number, row in enumerate(rows, start=2)}
    warnings = [warning for number, row in enumerate(rows, start=2) if not validation[number] for warning in row_warnings(context, payload.entity_type, row, number)]
    warnings.extend(within_file_warnings(payload.entity_type, rows))
    if warnings and not payload.confirm_warnings:
        raise http_error(409, "Review and confirm the possible duplicate warnings before importing.")
    errors: list[dict] = []
    imported = 0
    for number, row in enumerate(rows, start=2):
        found = validation[number]
        if found:
            errors.extend(found); continue
        owner = context["users"][row["owner_email"].casefold()]
        if payload.entity_type == "companies":
            db.add(Company(**stamp(user), name=row["name"], country=row["country"], owner_id=owner.id, domain=row["domain"].casefold(), industry=row["industry"], company_type=row["company_type"] or "Prospect", primary_region=row["primary_region"], global_account_name=row["global_account_name"]))
        elif payload.entity_type == "contacts":
            company = context["companies"][row["company_name"].casefold()]
            sourced_by = context["users"][row["sourced_by_email"].casefold()] if row["sourced_by_email"] else user
            dnc = row["do_not_contact"].casefold() in {"true", "yes", "1"}
            engagement = "Do not contact" if dnc else (row["engagement_status"] or "Not contacted")
            db.add(Contact(**stamp(user), company_id=company.id, first_name=row["first_name"], last_name=row["last_name"], country=row["country"], owner_id=owner.id, sourced_by_id=sourced_by.id, email=row["email"].casefold(), job_title=row["job_title"], seniority=row["seniority"] or "Unknown", phone=row["phone"], mobile=row["mobile"], linkedin_url=row["linkedin_url"], city=row["city"], source_channel=row["source_channel"] or "Other", source_detail=row["source_detail"], engagement_status=engagement, do_not_contact=dnc, consent_basis=row["consent_basis"] or "Business card or event", notes=row["notes"]))
        else:
            company = context["companies"][row["company_name"].casefold()]
            holder = context["users"][row["holder_email"].casefold()] if row["holder_email"] else owner
            sourced_by = context["users"][row["sourced_by_email"].casefold()] if row["sourced_by_email"] else user
            pursuit = Pursuit(**stamp(user), name=row["name"], company_id=company.id, owner_id=owner.id, sourced_by_id=sourced_by.id, holder_id=holder.id, next_action=row["next_action"], action_type=row["action_type"], action_date=date.fromisoformat(row["action_date"]), source_channel=row["source_channel"], source_detail=row["source_detail"], priority=row["priority"] or "Medium")
            db.add(pursuit); db.flush(); db.add(Lead(**stamp(user), pursuit_id=pursuit.id, area_of_interest=row["area_of_interest"]))
        db.flush(); imported += 1
    status = "Completed" if not errors else ("Failed" if not imported else "Completed with errors")
    job = ImportJob(**stamp(user), filename=payload.filename, file_type=payload.file_type, entity_type=payload.entity_type, status=status, total_rows=len(rows), imported_rows=imported, skipped_rows=len(rows) - imported, errors=errors[:500], warnings=warnings[:500])
    db.add(job); audit(db, user, "File import completed", detail=f"{payload.entity_type}: {imported}/{len(rows)} rows from {payload.filename}"); db.commit(); db.refresh(job)
    return present_job(job)


def present_job(job: ImportJob) -> dict:
    return {"id": str(job.id), "filename": job.filename, "file_type": job.file_type, "entity_type": job.entity_type, "status": job.status, "total_rows": job.total_rows, "imported_rows": job.imported_rows, "skipped_rows": job.skipped_rows, "errors": job.errors, "warnings": job.warnings, "created_at": job.created_at, "created_by_id": job.created_by_id}


@router.get("/imports/")
def import_history(db: Session = Depends(get_db), user: User = Depends(current_user)):
    jobs = db.scalars(scoped(db, ImportJob, user).order_by(ImportJob.created_at.desc()).limit(50)).all()
    return [present_job(job) for job in jobs]


@router.get("/search/")
def global_search(q: str = Query(min_length=2, max_length=120), entity_type: str = "all", owner_id: int | None = None, status_filter: str = "", country: str = "", page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), db: Session = Depends(get_db), user: User = Depends(current_user)):
    needle = f"%{q.strip()}%"; results: list[dict] = []
    if entity_type in {"all", "companies"}:
        statement = scoped(db, Company, user).where(or_(Company.name.ilike(needle), Company.domain.ilike(needle), Company.industry.ilike(needle))).options(selectinload(Company.owner))
        if owner_id: statement = statement.where(Company.owner_id == owner_id)
        if country: statement = statement.where(func.lower(Company.country) == country.casefold())
        for item in db.scalars(statement.limit(200)).all(): results.append({"type": "Company", "id": str(item.id), "title": item.name, "subtitle": item.domain or item.industry or item.country, "status": item.company_type, "owner": item.owner.display_name, "route": "companies"})
    if entity_type in {"all", "contacts"}:
        statement = scoped(db, Contact, user).where(or_(Contact.first_name.ilike(needle), Contact.last_name.ilike(needle), Contact.email.ilike(needle), Contact.phone.ilike(needle), Contact.mobile.ilike(needle))).options(selectinload(Contact.company), selectinload(Contact.owner))
        if owner_id: statement = statement.where(Contact.owner_id == owner_id)
        if country: statement = statement.where(func.lower(Contact.country) == country.casefold())
        for item in db.scalars(statement.limit(200)).all(): results.append({"type": "Contact", "id": str(item.id), "title": item.name, "subtitle": f"{item.company.name} · {item.email or item.phone or item.country}", "status": item.engagement_status, "owner": item.owner.display_name, "route": "contacts"})
    if entity_type in {"all", "leads", "opportunities"}:
        statement = scoped(db, Pursuit, user).where(or_(Pursuit.name.ilike(needle), Pursuit.source_detail.ilike(needle))).options(selectinload(Pursuit.company), selectinload(Pursuit.owner), selectinload(Pursuit.lead), selectinload(Pursuit.opportunity))
        if owner_id: statement = statement.where(Pursuit.owner_id == owner_id)
        for item in db.scalars(statement.limit(300)).all():
            kind = "Opportunity" if item.opportunity else "Lead"
            expected_type = "opportunities" if kind == "Opportunity" else "leads"
            if entity_type != "all" and entity_type != expected_type: continue
            state = item.opportunity.stage if item.opportunity else item.lead.status
            if status_filter and state != status_filter: continue
            if country and item.company.country.casefold() != country.casefold(): continue
            results.append({"type": kind, "id": str(item.id), "title": item.name, "subtitle": item.company.name, "status": state, "owner": item.owner.display_name, "route": "pipeline" if item.opportunity else "leads"})
    order = {"Opportunity": 0, "Lead": 1, "Company": 2, "Contact": 3}
    results.sort(key=lambda item: (order[item["type"]], item["title"].casefold()))
    total = len(results); start = (page - 1) * page_size
    return {"query": q, "total": total, "page": page, "page_size": page_size, "results": results[start:start + page_size]}


COMPANY_MERGE_FIELDS = ("name", "domain", "company_type", "industry", "country", "global_account_name", "primary_region", "owner_id")
CONTACT_MERGE_FIELDS = ("first_name", "last_name", "email", "job_title", "seniority", "phone", "mobile", "linkedin_url", "country", "city", "owner_id", "sourced_by_id", "source_channel", "source_detail", "engagement_status", "consent_basis", "notes")


def ordered_pair(first: UUID, second: UUID) -> tuple[UUID, UUID]:
    return tuple(sorted((first, second), key=str))


def company_record(item: Company) -> dict:
    return {"id": str(item.id), "name": item.name, "detail": item.domain or item.country, "fields": {"name": item.name, "domain": item.domain, "company_type": item.company_type, "industry": item.industry, "country": item.country, "global_account_name": item.global_account_name, "primary_region": item.primary_region, "owner_id": item.owner.display_name}}


def contact_record(item: Contact) -> dict:
    return {"id": str(item.id), "name": item.name, "detail": f"{item.company.name} · {item.email or item.phone or item.country}", "fields": {"first_name": item.first_name, "last_name": item.last_name, "email": item.email, "job_title": item.job_title, "seniority": item.seniority, "phone": item.phone, "mobile": item.mobile, "linkedin_url": item.linkedin_url, "country": item.country, "city": item.city, "owner_id": item.owner.display_name, "sourced_by_id": item.sourced_by.display_name, "source_channel": item.source_channel, "source_detail": item.source_detail, "engagement_status": item.engagement_status, "consent_basis": item.consent_basis, "notes": item.notes}}


def duplicate_groups(db: Session, user: User) -> list[dict]:
    companies = db.scalars(scoped(db, Company, user).options(selectinload(Company.owner)).order_by(Company.name)).all()
    contacts = db.scalars(scoped(db, Contact, user).options(selectinload(Contact.company), selectinload(Contact.owner), selectinload(Contact.sourced_by)).order_by(Contact.first_name, Contact.last_name)).all()
    dismissed = {ordered_pair(item.first_id, item.second_id) for item in db.scalars(scoped(db, DuplicateDecision, user).where(DuplicateDecision.decision == "Distinct")).all()}
    pairs: dict[tuple[str, UUID, UUID], dict] = {}

    def add(entity: str, first, second, reason: str, confidence: int):
        left, right = ordered_pair(first.id, second.id)
        if (left, right) in dismissed: return
        key = (entity, left, right)
        entry = pairs.setdefault(key, {"entity_type": entity, "match_reasons": [], "confidence": 0, "records": [company_record(first), company_record(second)] if entity == "companies" else [contact_record(first), contact_record(second)]})
        if reason not in entry["match_reasons"]: entry["match_reasons"].append(reason)
        entry["confidence"] = max(entry["confidence"], confidence)

    for attr, reason in (("name", "Exact company name"), ("domain", "Exact company domain")):
        buckets = defaultdict(list)
        for item in companies:
            key = normalized(getattr(item, attr));
            if key: buckets[key].append(item)
        for items in buckets.values():
            for first, second in combinations(items, 2): add("companies", first, second, reason, 100)
    company_blocks = defaultdict(list)
    for item in companies: company_blocks[(normalized(item.country), normalized(item.name)[:3])].append(item)
    for items in company_blocks.values():
        ordered = sorted(items, key=lambda item: normalized(item.name))
        for index, first in enumerate(ordered):
            for second in ordered[index + 1:index + 6]:
                score = similarity(first.name, second.name)
                if score >= 88: add("companies", first, second, "Similar company name", score)

    for attr, reason in (("email", "Exact email"), ("phone", "Exact phone")):
        buckets = defaultdict(list)
        for item in contacts:
            value = item.mobile or item.phone if attr == "phone" else item.email
            key = normalized(value)
            if key: buckets[key].append(item)
        for items in buckets.values():
            for first, second in combinations(items, 2): add("contacts", first, second, reason, 100)
    contact_blocks = defaultdict(list)
    for item in contacts: contact_blocks[(item.company_id, normalized(item.last_name)[:2])].append(item)
    for items in contact_blocks.values():
        ordered = sorted(items, key=lambda item: normalized(item.name))
        for index, first in enumerate(ordered):
            for second in ordered[index + 1:index + 6]:
                score = similarity(first.name, second.name)
                if score >= 90: add("contacts", first, second, "Similar contact name at same company", score)
    groups = list(pairs.values())
    for group in groups:
        fields = group["records"][0]["fields"]
        reasons = group["match_reasons"]
        if "Exact company domain" in reasons:
            group["match_type"] = "Company domain"
            group["match_value"] = normalized(fields.get("domain", ""))
        elif "Exact company name" in reasons:
            group["match_type"] = "Company name"
            group["match_value"] = normalized(fields.get("name", ""))
        elif "Exact email" in reasons:
            group["match_type"] = "Contact email"
            group["match_value"] = normalized(fields.get("email", ""))
        elif "Exact phone" in reasons:
            group["match_type"] = "Contact phone"
            group["match_value"] = normalized(fields.get("mobile") or fields.get("phone", ""))
        else:
            group["match_type"] = " · ".join(reasons)
            group["match_value"] = ":".join(sorted(record["id"] for record in group["records"]))
    return sorted(groups, key=lambda group: (-group["confidence"], group["entity_type"], group["records"][0]["name"].casefold()))


@router.get("/duplicates/")
def duplicates(db: Session = Depends(get_db), user: User = Depends(current_user)):
    groups = duplicate_groups(db, user)
    return {"group_count": len(groups), "record_count": sum(len(group["records"]) for group in groups), "groups": groups}


@router.post("/duplicates/{entity_type}/merge/")
def merge_duplicate(entity_type: str, payload: MergeInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    import_access(user)
    if payload.primary_id == payload.duplicate_id: raise http_error(422, "Choose two different records.")
    if entity_type == "companies":
        primary = db.scalar(scoped(db, Company, user).where(Company.id == payload.primary_id)); duplicate = db.scalar(scoped(db, Company, user).where(Company.id == payload.duplicate_id))
        if not primary or not duplicate: raise http_error(404, "Company not found.")
        allowed_fields = COMPANY_MERGE_FIELDS
        db.execute(update(Contact).where(Contact.tenant_id == user.tenant_id, Contact.company_id == duplicate.id).values(company_id=primary.id, updated_by_id=user.id))
        db.execute(update(Pursuit).where(Pursuit.tenant_id == user.tenant_id, Pursuit.company_id == duplicate.id).values(company_id=primary.id, updated_by_id=user.id))
        db.execute(update(Activity).where(Activity.tenant_id == user.tenant_id, Activity.company_id == duplicate.id).values(company_id=primary.id, updated_by_id=user.id))
        db.execute(update(PartnerInvolvement).where(PartnerInvolvement.tenant_id == user.tenant_id, PartnerInvolvement.company_id == duplicate.id).values(company_id=primary.id, updated_by_id=user.id))
    elif entity_type == "contacts":
        primary = db.scalar(scoped(db, Contact, user).where(Contact.id == payload.primary_id)); duplicate = db.scalar(scoped(db, Contact, user).where(Contact.id == payload.duplicate_id))
        if not primary or not duplicate: raise http_error(404, "Contact not found.")
        if primary.company_id != duplicate.company_id: raise http_error(422, "Merge the company records first, then review these contacts again.")
        allowed_fields = CONTACT_MERGE_FIELDS
        db.execute(update(Activity).where(Activity.tenant_id == user.tenant_id, Activity.contact_id == duplicate.id).values(contact_id=primary.id, updated_by_id=user.id))
        db.execute(update(Opportunity).where(Opportunity.tenant_id == user.tenant_id, Opportunity.primary_contact_id == duplicate.id).values(primary_contact_id=primary.id, updated_by_id=user.id))
        db.execute(update(PartnerInvolvement).where(PartnerInvolvement.tenant_id == user.tenant_id, PartnerInvolvement.contact_id == duplicate.id).values(contact_id=primary.id, updated_by_id=user.id))
        existing = set(db.scalars(select(PursuitContact.pursuit_id).where(PursuitContact.tenant_id == user.tenant_id, PursuitContact.contact_id == primary.id, PursuitContact.is_deleted.is_(False))).all())
        links = db.scalars(scoped(db, PursuitContact, user).where(PursuitContact.contact_id == duplicate.id)).all()
        for link in links:
            if link.pursuit_id in existing: link.is_deleted = True
            else: link.contact_id = primary.id
            link.updated_by_id = user.id
    else:
        raise http_error(404, "Merge type not found.")
    invalid_fields = set(payload.field_sources) - set(allowed_fields)
    if invalid_fields: raise http_error(422, {"field_sources": f'Unsupported merge fields: {", ".join(sorted(invalid_fields))}.'})
    valid_sources = {payload.primary_id, payload.duplicate_id}
    if any(source not in valid_sources for source in payload.field_sources.values()): raise http_error(422, {"field_sources": "Every field must come from one of the two reviewed records."})
    for field, source in payload.field_sources.items():
        if source == payload.duplicate_id: setattr(primary, field, getattr(duplicate, field))
    if entity_type == "contacts":
        primary.first_contacted_at = min((value for value in (primary.first_contacted_at, duplicate.first_contacted_at) if value), default=None)
        if primary.do_not_contact or duplicate.do_not_contact: primary.do_not_contact = True; primary.engagement_status = "Do not contact"
    duplicate.is_deleted = True; duplicate.updated_by_id = user.id
    audit(db, user, f"{entity_type[:-1].title()} records merged", detail=f"Kept {primary.id}; archived {duplicate.id}", before={"primary_id": str(primary.id), "duplicate_id": str(duplicate.id)}, after={"field_sources": {field: str(source) for field, source in payload.field_sources.items()}})
    db.commit()
    return {"detail": "Records merged.", "primary_id": str(primary.id), "archived_id": str(duplicate.id)}


@router.post("/duplicates/{entity_type}/dismiss/")
def dismiss_duplicate(entity_type: str, payload: DuplicateDismissInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    import_access(user)
    if entity_type not in {"companies", "contacts"}: raise http_error(404, "Duplicate type not found.")
    if payload.first_id == payload.second_id: raise http_error(422, "Choose two different records.")
    model = Company if entity_type == "companies" else Contact
    for identifier in (payload.first_id, payload.second_id):
        if not db.scalar(scoped(db, model, user).where(model.id == identifier)): raise http_error(404, "Record not found.")
    first, second = ordered_pair(payload.first_id, payload.second_id)
    item = db.scalar(scoped(db, DuplicateDecision, user).where(DuplicateDecision.entity_type == entity_type, DuplicateDecision.first_id == first, DuplicateDecision.second_id == second))
    if item:
        item.reason = payload.reason; item.reviewed_by_id = user.id; item.updated_by_id = user.id
    else:
        item = DuplicateDecision(**stamp(user), entity_type=entity_type, first_id=first, second_id=second, decision="Distinct", reason=payload.reason, reviewed_by_id=user.id); db.add(item)
    audit(db, user, "Duplicate suggestion dismissed", detail=f"{entity_type}: {first} / {second}; {payload.reason}"); db.commit()
    return {"detail": "The records were marked as distinct."}


@router.get("/quality/")
def data_quality(severity: str = "all", entity_type: str = "all", page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=200), db: Session = Depends(get_db), user: User = Depends(current_user)):
    issues: list[dict] = []
    companies = db.scalars(scoped(db, Company, user)).all()
    contacts = db.scalars(scoped(db, Contact, user).options(selectinload(Contact.company))).all()
    pursuits = db.scalars(scoped(db, Pursuit, user).options(selectinload(Pursuit.company), selectinload(Pursuit.lead), selectinload(Pursuit.opportunity))).all()
    def add(kind, item, code, reason, severity_name, route, action):
        issues.append({"type": kind, "id": str(item.id), "name": item.name, "code": code, "reason": reason, "severity": severity_name, "route": route, "recommended_action": action})
    for item in companies:
        if not item.domain: add("Company", item, "company_domain", "Missing website domain", "Medium", "companies", "Add the primary website domain.")
        if not item.industry: add("Company", item, "company_industry", "Missing industry", "Low", "companies", "Select or enter the industry.")
        if item.global_account_name and not item.primary_region: add("Company", item, "global_primary_region", "Global account has no primary region", "Medium", "companies", "Set the designated primary region.")
    for item in contacts:
        if not item.email: add("Contact", item, "contact_email", "Missing email address", "High", "contacts", "Add a verified business email if available.")
        if not item.phone and not item.mobile: add("Contact", item, "contact_phone", "Missing phone number", "Medium", "contacts", "Add a business phone or mobile number.")
        if not item.source_detail: add("Contact", item, "contact_source_detail", "Missing source detail", "Low", "contacts", "Record where or from whom the contact was obtained.")
        if item.do_not_contact and item.engagement_status != "Do not contact": add("Contact", item, "contact_dnc_status", "Do-not-contact flag and engagement status disagree", "High", "contacts", "Align the engagement status with the contact restriction.")
    today = date.today()
    for item in pursuits:
        closed = item.opportunity and item.opportunity.stage in {"won", "lost"} or item.lead and item.lead.status == "closed"
        route = "pipeline" if item.opportunity else "leads"
        if not closed and item.action_date < today: add("Pursuit", item, "overdue_action", "Next action is overdue", "High", route, "Complete or reschedule the next action.")
        if not closed and not item.next_action.strip(): add("Pursuit", item, "missing_action", "Missing next action", "High", route, "Set a clear next action and owner.")
    duplicates = duplicate_groups(db, user)
    possible_duplicate_records = sum(len(group["records"]) for group in duplicates)
    total = len(companies) + len(contacts) + len(pursuits)
    affected_records = {(issue["type"], issue["id"]) for issue in issues}
    for group in duplicates:
        kind = "Company" if group["entity_type"] == "companies" else "Contact"
        affected_records.update((kind, record["id"]) for record in group["records"])
    affected = len(affected_records)
    score = max(0, round(100 * (1 - affected / max(total, 1))))
    metrics = {"companies_missing_domain": sum(1 for i in issues if i["code"] == "company_domain"), "contacts_missing_email": sum(1 for i in issues if i["code"] == "contact_email"), "contacts_missing_phone": sum(1 for i in issues if i["code"] == "contact_phone"), "overdue_actions": sum(1 for i in issues if i["code"] == "overdue_action")}
    severity_counts = {name: sum(1 for item in issues if item["severity"] == name) for name in ("High", "Medium", "Low")}
    severity_order = {"High": 0, "Medium": 1, "Low": 2}; issues.sort(key=lambda item: (severity_order[item["severity"]], item["type"], item["name"].casefold()))
    filtered = [item for item in issues if (severity == "all" or item["severity"].casefold() == severity.casefold()) and (entity_type == "all" or item["type"].casefold() == entity_type.casefold())]
    start = (page - 1) * page_size; filtered_total = len(filtered)
    return {"score": score, "records_checked": total, "issue_count": len(issues), "filtered_issue_count": filtered_total, "duplicate_group_count": len(duplicates), "duplicate_record_count": possible_duplicate_records, "metrics": metrics, "severity_counts": severity_counts, "issues": filtered[start:start + page_size], "page": page, "page_size": page_size, "pages": max(1, (filtered_total + page_size - 1) // page_size), "checked_at": datetime.now(timezone.utc)}
